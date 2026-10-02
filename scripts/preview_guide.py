"""Conservative cell-layout check and previews; these are not Excel renders."""

from __future__ import annotations

import json
import math
from pathlib import Path
import re

import openpyxl
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / ".build"
FONT = Path(r"C:\Windows\Fonts\meiryo.ttc")
BOLD = Path(r"C:\Windows\Fonts\meiryob.ttc")


def font(cell):
    return ImageFont.truetype(str(BOLD if cell.font.bold else FONT), round(cell.font.sz * 4 / 3), index=2)


def lines(text, width, ft):
    result = []
    for para in str(text).split("\n"):
        line = ""
        # Keep short Latin words together. Long identifiers still wrap when needed.
        for token in re.findall(r"[A-Za-z0-9_.:/%=-]+|[^A-Za-z0-9_.:/%=-]", para):
            if line and ft.getlength(line + token) > width:
                result.append(line.rstrip())
                line = ""
            for char in token:
                if line and ft.getlength(line + char) > width:
                    result.append(line)
                    line = ""
                line += char
        result.append(line.rstrip())
    return result


def widths(ws):
    result = []
    for col in range(1, ws.max_column + 1):
        dim = next((v for v in ws.column_dimensions.values() if v.min <= col <= v.max), None)
        # XlsxWriter's 7-pixel width basis is conservative for Meiryo UI digits.
        result.append(round((dim.width if dim else 8.43) * 7 + 5))
    return result


def height(ws, row):
    return (ws.row_dimensions[row].height or ws.sheet_format.defaultRowHeight) * 4 / 3


def merged_bounds(ws, cell):
    for area in ws.merged_cells.ranges:
        if cell.coordinate in area:
            return area.min_col, area.min_row, area.max_col, area.max_row
    return cell.column, cell.row, cell.column, cell.row


def text_value(cell):
    return str(cell.value) if cell.value is not None else ""


def color(value, fallback):
    if value and value.type == "rgb":
        return "#" + value.rgb[-6:]
    return fallback


def inspect_layout(book):
    issues = []
    checked = 0
    for ws in book:
        cols = widths(ws)
        for row in ws:
            for cell in row:
                if cell.value is None or cell.data_type == "f":
                    continue
                c1, r1, c2, r2 = merged_bounds(ws, cell)
                ft = font(cell)
                wrapped = lines(text_value(cell), sum(cols[c1-1:c2]) - 12, ft)
                need = len(wrapped) * sum(ft.getmetrics()) + 8
                available = sum(height(ws, r) for r in range(r1, r2 + 1))
                checked += 1
                if need > available + 1:
                    issues.append(dict(sheet=ws.title, cell=cell.coordinate, required_pt=math.ceil(need * .75), available_pt=round(available*.75, 1), lines=len(wrapped)))
    return dict(check="Conservative font-metric estimate, not Excel rendering", checked_cells=checked, overflow_estimates=issues)


def preview(book, name, first, last):
    ws = book[name]
    cols = widths(ws)
    x = [0]
    for width in cols:
        x.append(x[-1] + width)
    selected = list(range(1, 6)) + list(range(first, last + 1)) if first > 6 else list(range(1, last + 1))
    selected = list(dict.fromkeys(selected))
    y = {}
    pos = 46
    for r in selected:
        y[r] = pos
        pos += round(height(ws, r))
    im = Image.new("RGB", (x[-1]+2, pos+2), "white")
    draw = ImageDraw.Draw(im)
    note_font = ImageFont.truetype(str(FONT), 17, index=2)
    draw.text((8, 9), f"配置の推定プレビュー（Excel実画面ではありません）  {name}", font=note_font, fill="#333333")
    for r in selected:
        for cell in ws[r]:
            if isinstance(cell, openpyxl.cell.cell.MergedCell):
                continue
            c1, r1, c2, r2 = merged_bounds(ws, cell)
            if r1 != r or r2 not in y:
                continue
            box = (x[c1-1], y[r], x[c2], y[r2]+round(height(ws, r2)))
            fill = color(cell.fill.fgColor, "#FFFFFF") if cell.fill.patternType == "solid" else "#FFFFFF"
            draw.rectangle(box, fill=fill, outline="#CCCCCC", width=1)
            if cell.value is None:
                continue
            ft = font(cell)
            wrapped = lines(text_value(cell), box[2]-box[0]-12, ft)
            lineheight = sum(ft.getmetrics())
            for i, text in enumerate(wrapped):
                draw.text((box[0]+6, box[1]+4+i*lineheight), text, font=ft, fill=color(cell.font.color, "#222222"))
    target = BUILD / f"layout_preview_{name[:2]}_{first}.png"
    im.save(target)
    return str(target)


def main():
    book = openpyxl.load_workbook(ROOT / "outputs" / "AWS_JMeter_EAP81_Observability_Guide.xlsx")
    report = inspect_layout(book)
    report["previews"] = [preview(book, *item) for item in [("00_案内", 6, 13), ("05_入口指標", 6, 11), ("08_JVM指標", 6, 13), ("09_メモリ深掘り", 6, 12), ("12_設定例", 6, 23), ("22_出典", 6, 10)]]
    (BUILD / "layout_validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
