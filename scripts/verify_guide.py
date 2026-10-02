"""Check document structure; optionally verify/render using installed Excel."""

from __future__ import annotations

import argparse
import ast
import gc
import hashlib
import json
import math
from pathlib import Path
import re
import zipfile

import openpyxl

from guide_content import CODE_EXAMPLES, SECTIONS, SOURCES

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / ".build"
OUT = ROOT / "outputs"
STEM = "AWS_JMeter_EAP81_Observability_Guide"


def evaluate_formula(formula, values):
    """Evaluate only the small arithmetic/IF subset used in this workbook."""
    def evaluate(node):
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.Name) and re.fullmatch(r"[A-Z]+[0-9]+", node.id):
            return values[node.id]
        if isinstance(node, ast.BinOp):
            left, right = evaluate(node.left), evaluate(node.right)
            if isinstance(node.op, ast.Div):
                return left / right
            if isinstance(node.op, ast.Mult):
                return left * right
        if isinstance(node, ast.Compare) and len(node.ops) == 1:
            left, right = evaluate(node.left), evaluate(node.comparators[0])
            if isinstance(node.ops[0], ast.Gt):
                return left > right
            if isinstance(node.ops[0], ast.GtE):
                return left >= right
            if isinstance(node.ops[0], ast.LtE):
                return left <= right
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
            if node.func.id == "IF" and len(node.args) == 3:
                return evaluate(node.args[1] if evaluate(node.args[0]) else node.args[2])
            if node.func.id == "AND":
                return all(evaluate(arg) for arg in node.args)
            if node.func.id == "NA" and not node.args:
                return math.nan
        raise AssertionError(f"Unsupported formula syntax: {ast.dump(node)}")

    return evaluate(ast.parse(formula.removeprefix("="), mode="eval").body)


def formula_verify(book, cached):
    ws = book["18_例と計算"]
    values = {cell.coordinate: cell.value for row in ws for cell in row if isinstance(cell.value, (int, float))}
    expected = [.5, 2 / 3, .05, 100, 100 / 6100, 200, 20]
    calculated = [evaluate_formula(ws.cell(r, 6).value, values) for r in range(6, 13)]
    stored = [cached["18_例と計算"].cell(r, 6).value for r in range(6, 13)]
    assert all(abs(a - b) < 1e-9 and abs(c - b) < 1e-9 for a, b, c in zip(calculated, expected, stored))
    cases = [("F6", "D6", 0), ("F6", "D6", -1), ("F7", "B7", 2000),
             ("F8", "B8", -1), ("F8", "D8", 0), ("F9", "D9", 0),
             ("F10", "B10", 7000), ("F11", "B11", -1), ("F12", "D12", -1)]
    for output, input_cell, bad_value in cases:
        assert math.isnan(evaluate_formula(ws[output].value, {**values, input_cell: bad_value})), (output, input_cell)
    return dict(method="Local restricted formula evaluator and cached-value comparison; not Excel recalculation", calculated=calculated, boundary_cases=len(cases))


def excel_verify():
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = None
    wb = None
    pdfs = []
    try:
        app = win32com.client.DispatchEx("Excel.Application")
        app.Visible = False
        app.DisplayAlerts = False
        app.AskToUpdateLinks = False
        app.AutomationSecurity = 3
        wb = app.Workbooks.Open(str(OUT / f"{STEM}.xlsx"), UpdateLinks=0, ReadOnly=False)
        # AutoFit only nonmerged table bodies. Keep deliberately laid-out drawings
        # and code cells intact. Add breathing space after actual Excel measurement.
        for section in SECTIONS:
            ws = wb.Worksheets(section["sheet"])
            last = 5 + len(section["rows"])
            ws.Range(f"A6:{openpyxl.utils.get_column_letter(len(section['headers']))}{last}").Rows.AutoFit()
            for row in range(6, last + 1):
                ws.Rows(row).RowHeight = min(390, max(50, ws.Rows(row).RowHeight + 10))
        ws = wb.Worksheets("22_出典")
        ws.Range(f"A6:E{len(SOURCES) + 5}").Rows.AutoFit()
        for row in range(6, len(SOURCES) + 6):
            ws.Rows(row).RowHeight = min(390, max(62, ws.Rows(row).RowHeight + 10))
        app.CalculateFullRebuild()
        result_sheet = wb.Worksheets("18_例と計算")
        actual = [result_sheet.Cells(row, 6).Value for row in range(6, 13)]
        expected = [.5, 2 / 3, .05, 100, 100 / 6100, 200, 20]
        assert all(abs(a - b) < 1e-9 for a, b in zip(actual, expected)), actual
        # Exercise boundary guards, then restore the illustrative workbook values.
        result_sheet.Cells(6, 4).Value = 0
        result_sheet.Cells(7, 2).Value = 2000
        result_sheet.Cells(9, 4).Value = 0
        app.CalculateFullRebuild()
        guards = [result_sheet.Cells(r, 6).Text for r in (6, 7, 9)]
        assert guards == ["#N/A"] * 3, guards
        result_sheet.Cells(6, 4).Value = 2048
        result_sheet.Cells(7, 2).Value = 1024
        result_sheet.Cells(9, 4).Value = 60
        app.CalculateFullRebuild()
        # Verify actual workbook font, not just the writer's requested style.
        font_checks = {}
        for name in ("00_案内", "05_入口指標", "08_JVM指標", "09_メモリ深掘り", "12_設定例"):
            font_checks[name] = wb.Worksheets(name).Cells(6, 1).Font.Name
        assert set(font_checks.values()) == {"Meiryo UI"}, font_checks
        wb.Worksheets("00_案内").Activate()
        wb.Save()
        for name in ("00_案内", "01_図解", "05_入口指標", "08_JVM指標", "09_メモリ深掘り", "12_設定例", "18_例と計算", "22_出典"):
            target = BUILD / f"preview_{name[:2]}.pdf"
            wb.Worksheets(name).ExportAsFixedFormat(0, str(target), 0, True, False)
            pdfs.append(str(target))
        return dict(excel_open=True, formula_values=actual, boundary_guards=guards, font_checks=font_checks, preview_pdfs=pdfs)
    finally:
        if wb is not None:
            wb.Close(SaveChanges=False)
        if app is not None:
            app.Quit()
        wb = None
        app = None
        gc.collect()
        pythoncom.CoUninitialize()


def structural_verify():
    path = OUT / f"{STEM}.xlsx"
    md_path = OUT / f"{STEM}.md"
    book = openpyxl.load_workbook(path, data_only=False)
    assert len(book.sheetnames) == 23, book.sheetnames
    assert book.sheetnames == sorted(book.sheetnames)
    fonts = set()
    links = 0
    nonempty = 0
    issues = []
    for ws in book:
        assert ws.freeze_panes, ws.title
        assert ws.sheet_view.showGridLines is False, ws.title
        assert ws.print_area, ws.title
        for row in ws:
            for cell in row:
                if cell.value is None:
                    continue
                nonempty += 1
                fonts.add(cell.font.name)
                assert cell.alignment.wrap_text, (ws.title, cell.coordinate, "wrap disabled")
                assert cell.data_type != "e", (ws.title, cell.coordinate, cell.value)
                if isinstance(cell.value, str):
                    assert "\ufffd" not in cell.value
                    assert len(cell.value) <= 32767
                if cell.hyperlink:
                    links += 1
                    location = cell.hyperlink.location
                    if location:
                        match = re.match(r"'?([^']+)'?!([A-Z]+[0-9]+)$", location)
                        assert match, location
                        assert match[1] in book.sheetnames, location
    assert fonts == {"Meiryo UI"}, fonts
    for section in SECTIONS:
        ws = book[section["sheet"]]
        assert ws.auto_filter.ref
        for r, expected in enumerate(section["rows"], 6):
            actual = tuple(ws.cell(r, c + 1).value for c in range(len(expected) - 1))
            assert actual == tuple(expected[:-1]), (section["sheet"], r)
            if ws.row_dimensions[r].height and ws.row_dimensions[r].height >= 390:
                issues.append((ws.title, r, "near Excel row-height cap"))
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        charts = [n for n in archive.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n)]
        media = [n for n in archive.namelist() if n.startswith("xl/media/")]
        assert len(charts) == 2, charts
        assert len(media) == 3, media
    md = md_path.read_text(encoding="utf-8")
    assert md.count("```mermaid") == 3
    assert len(re.findall(r"^```", md, re.M)) % 2 == 0
    assert "\ufffd" not in md
    assert all(line == line.rstrip() for line in md.splitlines()), "trailing whitespace"
    for key, s in SOURCES.items():
        assert s["url"] in md, key
    for example in CODE_EXAMPLES:
        if example["language"] == "json":
            json.loads(example["code"])
    data = openpyxl.load_workbook(path, data_only=True)
    values = [data["18_例と計算"].cell(r, 6).value for r in range(6, 13)]
    assert all(isinstance(v, (float, int)) for v in values), values
    formulas = formula_verify(book, data)
    assert not issues, issues
    return dict(sheets=len(book.sheetnames), nonempty_cells=nonempty, fonts=sorted(fonts), hyperlinks=links,
                charts=len(charts), diagrams=len(media), sources=len(SOURCES), code_examples=len(CODE_EXAMPLES),
                markdown_chars=len(md), xlsx_bytes=path.stat().st_size, markdown_bytes=md_path.stat().st_size,
                formula_cached_values=values, formula_validation=formulas,
                sha256={p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (path, md_path)})


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", action="store_true")
    args = parser.parse_args()
    result = {}
    if args.excel:
        result["excel"] = excel_verify()
    result["structure"] = structural_verify()
    (BUILD / "validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
